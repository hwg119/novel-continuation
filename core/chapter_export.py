# -*- coding: utf-8 -*-
"""把章节正文及其已保存插图导出为 Word 文档。"""
import re
import xml.etree.ElementTree as ET
from pathlib import Path


WIDTH, HEIGHT = 1200, 520
SVG_NS = "http://www.w3.org/2000/svg"
ET.register_namespace("", SVG_NS)


def render_svg_png(svg: str, png_path: str, scale: int = 2) -> None:
    """统一使用 resvg，保持正文预览与 Word 插图一致。"""
    from core.svg_renderer import render_resvg_png
    render_resvg_png(svg, png_path, scale)


def export_chapter_docx(title: str, body: str, output_path: str,
                        chapter_number: int | None = None,
                        illustration_svg: str | None = None) -> dict:
    """导出 Word；仅在提供已保存插图时才在标题后插图。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    def set_chinese_font(style, family):
        """显式设置东亚字体，避免 Word 用西文字体替代中文。"""
        style.font.name = family
        properties = style._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{kind}"), family)

    title = (title or "").strip() or "未命名章节"
    if chapter_number is not None:
        if chapter_number < 1:
            raise ValueError("章号必须大于零")
        chapter_label = f"第 {chapter_number} 章"
        display_title = title if re.match(r"^第\s*\d+\s*章", title) else f"{chapter_label}  {title}"
    else:
        display_title = title
    body = (body or "").strip()
    if not body:
        raise ValueError("当前章正文为空")
    path = Path(output_path)
    if path.suffix.lower() != ".docx":
        raise ValueError("导出路径必须以 .docx 结尾")
    path.parent.mkdir(parents=True, exist_ok=True)
    png_path = path.with_suffix(".png") if illustration_svg else None
    if illustration_svg and png_path:
        from core.chapter_illustration import validate_illustration_svg
        render_svg_png(validate_illustration_svg(illustration_svg), str(png_path))

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    set_chinese_font(normal, "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(40, 43, 46)
    # 正文使用较紧凑的行距；自然段之间仍由段后间距保持层次。
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    title_style = doc.styles["Title"]
    set_chinese_font(title_style, "Microsoft YaHei")
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    heading = doc.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(12)
    run = heading.add_run(display_title)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    if png_path:
        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_after = Pt(18)
        picture.add_run().add_picture(str(png_path), width=Inches(6.6))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("第 ")
    footer_run.font.name = "Microsoft YaHei"
    footer_run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer_run = footer.add_run(" 页")
    footer_run.font.name = "Microsoft YaHei"
    footer_run.font.size = Pt(9)
    for block in re.split(r"\n\s*\n", body):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        for index, line in enumerate(lines):
            para = doc.add_paragraph(line.strip())
            para.paragraph_format.first_line_indent = Pt(25 if index == 0 else 0)
            para.paragraph_format.space_after = Pt(6 if index == len(lines) - 1 else 0)
    doc.save(str(path))
    return {"docx": str(path), "png": str(png_path) if png_path else None,
            "illustrated": bool(png_path)}


def export_chapters_docx(chapters: list[dict], output_path: str) -> dict:
    """把多个章节按章分页合并到一个 Word 文档中。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    if not chapters:
        raise ValueError("请至少选择一个章节")
    path = Path(output_path)
    if path.suffix.lower() != ".docx":
        raise ValueError("导出路径必须以 .docx 结尾")
    path.parent.mkdir(parents=True, exist_ok=True)

    def set_chinese_font(style, family):
        style.font.name = family
        properties = style._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{kind}"), family)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    set_chinese_font(normal, "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(40, 43, 46)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    set_chinese_font(doc.styles["Title"], "Microsoft YaHei")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)

    image_paths = []
    for chapter_index, chapter in enumerate(chapters):
        number = int(chapter.get("number") or 0)
        if number < 1:
            raise ValueError("章号必须大于零")
        title = str(chapter.get("title") or "").strip() or "未命名章节"
        body = str(chapter.get("body") or "").strip()
        if not body:
            raise ValueError(f"第 {number} 章正文为空")
        if chapter_index:
            doc.add_page_break()
        display_title = title if re.match(r"^第\s*\d+\s*章", title) else f"第 {number} 章  {title}"
        heading = doc.add_paragraph(style="Title")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.space_after = Pt(12)
        run = heading.add_run(display_title)
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        svg = chapter.get("illustration_svg")
        if svg:
            from core.chapter_illustration import validate_illustration_svg
            png_path = path.with_name(f"{path.stem}_chapter_{number}.png")
            render_svg_png(validate_illustration_svg(str(svg)), str(png_path))
            image_paths.append(str(png_path))
            picture = doc.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.space_after = Pt(18)
            picture.add_run().add_picture(str(png_path), width=Inches(6.6))
        for block in re.split(r"\n\s*\n", body):
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            for index, line in enumerate(lines):
                para = doc.add_paragraph(line)
                para.paragraph_format.first_line_indent = Pt(25 if index == 0 else 0)
                para.paragraph_format.space_after = Pt(6 if index == len(lines) - 1 else 0)
    doc.save(str(path))
    return {"docx": str(path), "pngs": image_paths, "chapters": len(chapters)}


def export_learning_docx(title: str, edition: dict, output_path: str,
                         chapter_number: int, mode: str = "inline",
                         include_vocabulary: bool = True,
                         illustration_svg: str | None = None) -> dict:
    """把已生成的分级学习版导出为适合打印的 Word。"""
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    if mode not in ("inline", "endnotes", "exercise"):
        raise ValueError("学习版 Word 模式无效")
    if not edition.get("valid") or not edition.get("blocks"):
        raise ValueError("当前英语学习版不存在或已经过期")
    path = Path(output_path)
    if path.suffix.lower() != ".docx":
        raise ValueError("导出路径必须以 .docx 结尾")
    path.parent.mkdir(parents=True, exist_ok=True)
    png_path = path.with_suffix(".png") if illustration_svg else None
    if illustration_svg and png_path:
        from core.chapter_illustration import validate_illustration_svg
        render_svg_png(validate_illustration_svg(illustration_svg), str(png_path))

    def set_font(target, family: str):
        target.font.name = family
        properties = target._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{kind}"), family)

    def set_cell_shading(cell, fill: str):
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    def set_cell_margins(cell, top=45, start=70, bottom=45, end=70):
        properties = cell._tc.get_or_add_tcPr()
        margins = properties.first_child_found_in("w:tcMar")
        if margins is None:
            margins = OxmlElement("w:tcMar")
            properties.append(margins)
        for name, value in (("top", top), ("start", start),
                            ("bottom", bottom), ("end", end)):
            node = margins.find(qn(f"w:{name}"))
            if node is None:
                node = OxmlElement(f"w:{name}")
                margins.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def set_table_borders(table, color="D9D9D9", size="4"):
        properties = table._tbl.tblPr
        borders = properties.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            properties.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = borders.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                borders.append(node)
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), size)
            node.set(qn("w:color"), color)

    def set_fixed_table_widths(table, widths):
        """同时写入表格网格和单元格宽度，避免 Word 自动重排窄列。"""
        width_twips = [round(width.inches * 1440) for width in widths]
        properties = table._tbl.tblPr
        layout = properties.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            properties.append(layout)
        layout.set(qn("w:type"), "fixed")
        table_width = properties.first_child_found_in("w:tblW")
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            properties.append(table_width)
        table_width.set(qn("w:w"), str(sum(width_twips)))
        table_width.set(qn("w:type"), "dxa")

        grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
        for grid_column, width in zip(grid_columns, width_twips):
            grid_column.set(qn("w:w"), str(width))
        for row in table.rows:
            for cell, width in zip(row.cells, width_twips):
                cell_properties = cell._tc.get_or_add_tcPr()
                cell_width = cell_properties.first_child_found_in("w:tcW")
                if cell_width is None:
                    cell_width = OxmlElement("w:tcW")
                    cell_properties.append(cell_width)
                cell_width.set(qn("w:w"), str(width))
                cell_width.set(qn("w:type"), "dxa")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.82)
    normal = doc.styles["Normal"]
    set_font(normal, "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(38, 43, 46)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    title_style = doc.styles["Title"]
    set_font(title_style, "Microsoft YaHei")
    title_style.font.color.rgb = RGBColor(0, 0, 0)

    display_title = (title or "未命名章节").strip()
    if not re.match(r"^第\s*\d+\s*章", display_title):
        display_title = f"第 {chapter_number} 章  {display_title}"
    heading = doc.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run(display_title)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(13)
    subtitle_run = subtitle.add_run(f"英语学习版  {edition.get('level_label', '')}")
    subtitle_run.font.name = "Microsoft YaHei"
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(82, 92, 91)

    if png_path:
        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_after = Pt(15)
        picture.add_run().add_picture(str(png_path), width=Inches(6.45))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)

    references = []
    vocabulary = {}
    for block in edition["blocks"]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.space_after = Pt(6)
        inline_sources = []
        for segment in block.get("segments", []):
            if segment.get("translated"):
                number = len(references) + 1
                translation = (segment.get("translation") if mode == "exercise"
                               else segment.get("text")) or segment.get("source", "")
                english = paragraph.add_run(translation)
                english.font.name = "Aptos"
                english.font.color.rgb = RGBColor(25, 63, 78)
                english.font.bold = False
                if mode in ("endnotes", "exercise"):
                    marker = paragraph.add_run(f"[{number}]")
                    marker.font.superscript = True
                    marker.font.size = Pt(8)
                    marker.font.color.rgb = RGBColor(125, 94, 38)
                references.append({"number": number, "source": segment.get("source", ""),
                                   "translation": segment.get("translation", "")})
                if mode == "inline":
                    inline_sources.append((number, segment.get("source", "")))
                for note in segment.get("annotations") or []:
                    key = str(note.get("word") or "").strip().lower()
                    if key and key not in vocabulary:
                        vocabulary[key] = {"word": note.get("word"),
                                           "meaning": note.get("meaning")}
            else:
                paragraph.add_run(segment.get("text") or segment.get("source", ""))
        if inline_sources:
            for number, source in inline_sources:
                source_para = doc.add_paragraph()
                source_para.paragraph_format.left_indent = Pt(24)
                source_para.paragraph_format.right_indent = Pt(12)
                source_para.paragraph_format.space_after = Pt(4)
                source_run = source_para.add_run(f"原句 {number}  {source}")
                source_run.font.name = "Microsoft YaHei"
                source_run.font.size = Pt(9.5)
                source_run.font.color.rgb = RGBColor(105, 119, 116)

    if mode == "endnotes" and references:
        doc.add_page_break()
        heading = doc.add_heading("章末原句索引", level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
        index_table = doc.add_table(rows=1, cols=3)
        index_table.autofit = False
        widths = (Inches(0.4), Inches(3.1), Inches(3.0))
        headers = index_table.rows[0].cells
        labels = ("编号", "英文句", "中文原句")
        for cell, label, width in zip(headers, labels, widths):
            cell.width = width
            cell.text = label
            set_cell_shading(cell, "E5EEEA")
            set_cell_margins(cell, top=55, bottom=55)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            for run in para.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
        row_properties = index_table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        row_properties.append(repeat)
        for item in references:
            cells = index_table.add_row().cells
            values = (str(item["number"]), item["translation"], item["source"])
            for column, (cell, value, width) in enumerate(zip(cells, values, widths)):
                cell.width = width
                cell.text = value
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if item["number"] % 2 == 0:
                    set_cell_shading(cell, "F7F9F8")
                para = cell.paragraphs[0]
                para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if column == 0
                                  else WD_ALIGN_PARAGRAPH.LEFT)
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.name = "Aptos" if column == 1 else "Microsoft YaHei"
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(47, 55, 55)
        set_fixed_table_widths(index_table, widths)
        set_table_borders(index_table)
    elif mode == "exercise" and references:
        doc.add_page_break()
        heading = doc.add_heading("参考原句", level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
        for item in references:
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(5)
            lead = para.add_run(f"{item['number']}  ")
            lead.bold = True
            lead.font.color.rgb = RGBColor(44, 91, 79)
            para.add_run(item["source"])
            english = doc.add_paragraph(item["translation"])
            english.paragraph_format.left_indent = Pt(22)
            english.paragraph_format.space_after = Pt(8)
            for run in english.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(91, 107, 105)

    if include_vocabulary and vocabulary:
        doc.add_heading("本章词汇", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(2.3)
        table.columns[1].width = Inches(4.2)
        headers = table.rows[0].cells
        headers[0].text, headers[1].text = "Word", "中文释义"
        for cell in headers:
            set_cell_shading(cell, "355E59")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        for index, item in enumerate(vocabulary.values()):
            cells = table.add_row().cells
            cells[0].text = str(item["word"] or "")
            cells[1].text = str(item["meaning"] or "")
            if index % 2:
                for cell in cells:
                    set_cell_shading(cell, "F1F5F3")
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.space_before = Pt(3)
        table.style = "Table Grid"

    doc.save(str(path))
    return {"docx": str(path), "png": str(png_path) if png_path else None,
            "illustrated": bool(png_path), "mode": mode,
            "vocabulary": len(vocabulary) if include_vocabulary else 0}
